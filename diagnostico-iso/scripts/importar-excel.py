#!/usr/bin/env python3
"""Importa un diagnóstico diligenciado en Excel (formato del instrumento HUC) a un archivo JSON
compatible con server/seed/caso-*.json, o lo carga directamente por la API.

Uso:
  python3 scripts/importar-excel.py archivo.xlsx --organizacion "Nombre" --sigla XX --salida caso.json
  python3 scripts/importar-excel.py archivo.xlsx --api https://diagnosticoiso.proyectoia.org --email admin@... --password ... --org-id <uuid>

Requiere: pip install openpyxl requests
"""
import argparse, json, sys
try:
    import openpyxl
except ImportError:
    sys.exit('Instale openpyxl: pip install openpyxl')

MAPEO = {
    'cumple': 'cumple', 'evidencia completa': 'cumple',
    'evidencia parcial': 'cumple_parcial', 'cumplimiento parcial': 'cumple_parcial', 'parcial': 'cumple_parcial',
    'no cumple': 'no_cumple', 'brecha documental': 'no_cumple',
    'sin evidencia': 'sin_evidencia', 'sin evidencia aportada': 'sin_evidencia', 'antecedente por verificar': 'sin_evidencia',
}

def leer(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.worksheets[0]
    fila_enc = None
    for r in ws.iter_rows(min_row=1, max_row=15):
        vals = [str(c.value).strip().lower() if c.value else '' for c in r]
        if 'id' in vals and any('pregunta' in v for v in vals):
            fila_enc = r[0].row; enc = vals; break
    if not fila_enc:
        sys.exit('No se encontró la fila de encabezados (ID, Capítulo, Numeral ISO, Pregunta, Estado, Evidencia, Observaciones, Responsable).')
    col = lambda nombre: next((i for i, v in enumerate(enc) if nombre in v), None)
    ci, ce, cev, cob, cre = col('id'), col('estado'), col('evidencia'), col('observ'), col('responsable')
    out = []
    for r in ws.iter_rows(min_row=fila_enc + 1, values_only=True):
        if not r[ci]: continue
        estado = str(r[ce] or '').strip()
        val = MAPEO.get(estado.lower())
        out.append({'codigo': str(r[ci]).strip(), 'valoracion': val, 'estado_origen': estado if estado and estado.lower() not in ('cumple', 'no cumple', 'cumplimiento parcial', 'sin evidencia') else None,
                    'evidencia': str(r[cev] or '').strip() if cev is not None else '', 'observaciones': str(r[cob] or '').strip() if cob is not None else '', 'responsable': str(r[cre] or '').strip() if cre is not None else ''})
    return out

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('xlsx'); ap.add_argument('--organizacion', default='Organización'); ap.add_argument('--sigla', default=None)
    ap.add_argument('--titulo', default='Diagnóstico importado desde Excel'); ap.add_argument('--fecha', default=None)
    ap.add_argument('--salida', default=None)
    ap.add_argument('--api'); ap.add_argument('--email'); ap.add_argument('--password'); ap.add_argument('--org-id')
    a = ap.parse_args()
    resp = leer(a.xlsx)
    print(f'{len(resp)} respuestas leídas; sin mapear: {sum(1 for r in resp if r["valoracion"] is None)}', file=sys.stderr)
    if a.api:
        import requests
        s = requests.Session()
        s.post(a.api + '/api/auth/local/login', json={'email': a.email, 'password': a.password}).raise_for_status()
        if not a.org_id:
            r = s.post(a.api + '/api/organizaciones', json={'nombre': a.organizacion, 'sigla': a.sigla}); r.raise_for_status(); a.org_id = r.json()['id']
        r = s.post(f'{a.api}/api/organizaciones/{a.org_id}/diagnosticos', json={'titulo': a.titulo, 'fecha': a.fecha}); r.raise_for_status(); diag = r.json()['id']
        items = {i['codigo']: i['id'] for i in s.get(f'{a.api}/api/diagnosticos/{diag}').json()['items']}
        for x in resp:
            if x['codigo'] in items:
                s.put(f'{a.api}/api/diagnosticos/{diag}/respuestas/{items[x["codigo"]]}', json={k: x[k] for k in ('valoracion', 'evidencia', 'observaciones', 'responsable')}).raise_for_status()
        print(f'Diagnóstico {diag} creado en la organización {a.org_id}', file=sys.stderr)
    else:
        caso = {'organizacion': {'nombre': a.organizacion, 'sigla': a.sigla}, 'diagnostico': {'titulo': a.titulo, 'fecha': a.fecha, 'estado': 'cerrado', 'instrumento': 'iso9001-2015-amd1-2024'}, 'respuestas': resp}
        txt = json.dumps(caso, ensure_ascii=False, indent=1)
        if a.salida: open(a.salida, 'w', encoding='utf-8').write(txt)
        else: print(txt)

if __name__ == '__main__':
    main()
